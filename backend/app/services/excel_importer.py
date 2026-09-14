import io
import re
import uuid
from typing import Dict, Any, List, Optional, Tuple
from datetime import datetime, timezone
import pandas as pd
import openpyxl

from app.db.connection import db
from app.db.utils import id_query, parse_id
from app.models.canonical_models import (
    CanonicalStudent, IdentityData, ContactData, AcademicIdentityData,
    EducationData, ScoreRecord, UndergraduateEducation, EligibilityData,
    EntranceExamRecord, LinksData, DocumentsData, AuthenticationState,
    ApprovalState, ImportMetadata
)
from app.services.score_normalizer import normalize_academic_score
from app.services.board_conversion_service import (
    find_matching_conversion_rule, apply_conversion_rule, ensure_default_board_rules
)
from app.services.academic_timeline_service import (
    get_student_academic_status, get_or_create_academic_batch, get_or_create_class_section,
    ensure_default_academic_calendar
)

STANDARD_TEMPLATE_HEADERS = [
    "Student ID / Roll Number", "Full Name", "First Name", "Middle Name", "Last Name",
    "Date of Birth", "Gender", "Institutional Email", "Personal Email", "Mobile Number",
    "Program", "Department", "Class / Section", "Year of Passing",
    "Admission Year", "Expected Graduation Year", "Academic Batch",
    "10th Board", "10th School / Institution", "10th Passing Year", "10th Score Type", "10th Score",
    "12th / Diploma Board", "12th / Diploma Institution", "12th / Diploma Passing Year", "12th / Diploma Score Type", "12th / Diploma Score",
    "UG College / Institution", "UG Overall Score Type", "UG Overall Score",
    "Semester 1 SGPA", "Semester 2 SGPA", "Semester 3 SGPA", "Semester 4 SGPA",
    "Semester 5 SGPA", "Semester 6 SGPA", "Semester 7 SGPA", "Semester 8 SGPA",
    "Current Backlogs", "Backlog Count", "Education Gap", "First Attempt Status",
    "Entrance Exam", "Entrance Rank", "Nationality",
    "Profile Photo URL", "Combined Document URL", "Technical Certificates URL", "Achievement Certificates URL"
]

def clean_cell_str(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, (datetime, pd.Timestamp)):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    if s.lower() in ["nan", "none", "null", "undefined", "nat"]:
        return ""
    if s.endswith(".0") and not re.search(r"\.\d{2,}", s):
        try:
            s = str(int(float(s)))
        except ValueError:
            pass
    return s

def clean_email(val: Any) -> str:
    s = clean_cell_str(val).lower()
    match = re.search(r"[\w\.-]+@[\w\.-]+\.\w+", s)
    return match.group(0) if match else s

def clean_phone(val: Any) -> str:
    s = clean_cell_str(val)
    digits = re.sub(r"\D", "", s)
    if len(digits) == 12 and digits.startswith("91"):
        digits = digits[2:]
    return digits

async def resolve_hierarchy_batch_and_faculty(
    institution_id: str,
    program_name: str,
    dept_name: str,
    class_name: str,
    admission_year_val: Optional[Any],
    graduation_year_val: Optional[Any],
    batch_label_val: Optional[str],
    all_programs: List[Dict[str, Any]],
    all_depts: List[Dict[str, Any]],
    all_classes: List[Dict[str, Any]],
    all_faculty: List[Dict[str, Any]]
) -> Dict[str, Any]:
    """
    Resolves Program -> Duration -> Admission & Graduation Years -> Academic Batch -> Department -> Batch-specific Class -> Faculty Coordinator.
    """
    norm_prog = program_name.lower().strip() if program_name else ""
    norm_dept = dept_name.lower().strip() if dept_name else ""
    norm_cls = class_name.lower().strip() if class_name else ""

    # 1. Resolve Program
    matched_prog = None
    if norm_prog:
        matched_prog = next((p for p in all_programs if (p.get("program_name") or p.get("name", "")).lower() == norm_prog or norm_prog in (p.get("program_name") or p.get("name", "")).lower()), None)
    if not matched_prog and all_programs:
        matched_prog = all_programs[0]

    prog_id = str(matched_prog.get("id") or matched_prog["_id"]) if matched_prog else None
    resolved_prog_name = matched_prog.get("program_name") or matched_prog.get("name", program_name) if matched_prog else program_name
    duration_years = int(matched_prog.get("duration_years") or 4) if matched_prog else 4

    # 2. Determine and Validate Admission & Expected Graduation Years
    grad_yr_int = None
    adm_yr_int = None
    batch_duration_mismatch = False

    if graduation_year_val:
        try:
            grad_yr_int = int(re.sub(r"\D", "", str(graduation_year_val))[:4])
        except Exception:
            grad_yr_int = None

    if admission_year_val:
        try:
            adm_yr_int = int(re.sub(r"\D", "", str(admission_year_val))[:4])
        except Exception:
            adm_yr_int = None

    if grad_yr_int and not adm_yr_int:
        adm_yr_int = grad_yr_int - duration_years
    elif adm_yr_int and not grad_yr_int:
        grad_yr_int = adm_yr_int + duration_years
    elif adm_yr_int and grad_yr_int:
        if (grad_yr_int - adm_yr_int) != duration_years:
            batch_duration_mismatch = True

    if not grad_yr_int:
        grad_yr_int = datetime.now().year + (duration_years if datetime.now().month < 6 else duration_years - 1)
        adm_yr_int = grad_yr_int - duration_years

    batch_label = batch_label_val.strip() if batch_label_val else f"{adm_yr_int}-{grad_yr_int}"

    # 3. Resolve / Create Academic Batch
    academic_batch_id = None
    if prog_id and adm_yr_int and grad_yr_int:
        batch_doc = await get_or_create_academic_batch(institution_id, prog_id, adm_yr_int, grad_yr_int)
        academic_batch_id = batch_doc.get("id") or str(batch_doc["_id"])
        batch_label = batch_doc.get("batch_label") or batch_label

    # 4. Resolve Department
    matched_dept = None
    if norm_dept:
        matched_dept = next((d for d in all_depts if (d.get("department_name") or d.get("name", "")).lower() == norm_dept or norm_dept in (d.get("department_name") or d.get("name", "")).lower()), None)
    if not matched_dept and all_depts:
        matched_dept = all_depts[0]

    dept_id = str(matched_dept.get("id") or matched_dept["_id"]) if matched_dept else None
    resolved_dept_name = matched_dept.get("department_name") or matched_dept.get("name", dept_name) if matched_dept else dept_name

    # 5. Resolve / Create Batch-Specific Class Section
    class_id = None
    resolved_class_name = class_name
    faculty_id = None
    faculty_name = None

    if prog_id and dept_id and academic_batch_id:
        cls_doc = await get_or_create_class_section(
            institution_id, prog_id, dept_id, academic_batch_id, norm_cls or "Section-A"
        )
        class_id = cls_doc.get("id") or str(cls_doc["_id"])
        resolved_class_name = cls_doc.get("display_name") or cls_doc.get("class_name", class_name)

        # Resolve Faculty Coordinator mapped to this batch's class
        assigned_faculty = None
        if cls_doc.get("faculty_coordinator_id"):
            assigned_faculty = await db.db.users.find_one(id_query(cls_doc["faculty_coordinator_id"]))
        if not assigned_faculty:
            assigned_faculty = next((f for f in all_faculty if str(f.get("class_id")) == class_id or parse_id(str(f.get("class_id"))) == parse_id(class_id)), None)
        if not assigned_faculty and dept_id:
            assigned_faculty = next((f for f in all_faculty if str(f.get("department_id")) == dept_id), None)

        faculty_id = str(assigned_faculty["_id"]) if assigned_faculty else None
        faculty_name = assigned_faculty.get("name") if assigned_faculty else None

    is_mapped = bool(prog_id and dept_id and academic_batch_id and class_id and faculty_id and not batch_duration_mismatch)

    return {
        "program_id": prog_id,
        "program_name": resolved_prog_name,
        "duration_years": duration_years,
        "admission_year": adm_yr_int,
        "expected_graduation_year": grad_yr_int,
        "academic_batch_id": academic_batch_id,
        "batch_label": batch_label,
        "batch_duration_mismatch": batch_duration_mismatch,
        "department_id": dept_id,
        "department_name": resolved_dept_name,
        "class_id": class_id,
        "class_name": resolved_class_name,
        "faculty_id": faculty_id,
        "faculty_name": faculty_name,
        "is_mapped": is_mapped
    }

async def process_excel_import(
    file_bytes: bytes,
    institution_id: str,
    source_filename: str = "upload.xlsx",
    confirm_save: bool = False,
    imported_by_user_id: Optional[str] = None
) -> Dict[str, Any]:
    """
    Parses and validates standard template Excel with Academic Batch, Normalization, and Faculty Assignment.
    """
    await ensure_default_board_rules()
    await ensure_default_academic_calendar(institution_id)

    # Load workbook using openpyxl with pandas fallback
    header_row = []
    rows = []
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        if "Student_Master_Upload" in wb.sheetnames:
            ws = wb["Student_Master_Upload"]
        else:
            ws = wb.active
            
        data = list(ws.iter_rows(values_only=True))
        if not data or len(data) < 2:
            return {"success": False, "error": "Excel workbook is empty or contains no data rows."}
            
        header_row = [clean_cell_str(c) for c in data[0]]
        rows = data[1:]
    except Exception as e:
        # Fallback to pandas read_excel / read_csv
        try:
            df = pd.read_excel(io.BytesIO(file_bytes))
            if df.empty:
                return {"success": False, "error": "Spreadsheet contains no data rows."}
            header_row = [clean_cell_str(c) for c in df.columns]
            rows = [tuple(r) for r in df.values]
        except Exception:
            try:
                df = pd.read_csv(io.BytesIO(file_bytes))
                if df.empty:
                    return {"success": False, "error": "CSV file contains no data rows."}
                header_row = [clean_cell_str(c) for c in df.columns]
                rows = [tuple(r) for r in df.values]
            except Exception as e_inner:
                return {"success": False, "error": f"Failed to parse Excel file: {str(e)}"}

    # Load existing hierarchy from MongoDB
    all_programs = await db.db.programs.find({"$or": [{"institution_id": institution_id}, {"institution_id": parse_id(institution_id)}]}).to_list(length=200)
    all_depts = await db.db.departments.find({"$or": [{"institution_id": institution_id}, {"institution_id": parse_id(institution_id)}]}).to_list(length=500)
    all_classes = await db.db.classes.find({"$or": [{"institution_id": institution_id}, {"institution_id": parse_id(institution_id)}]}).to_list(length=1000)
    all_faculty = await db.db.users.find({"institution_id": institution_id, "role": "Faculty", "is_active": {"$ne": False}}).to_list(length=500)

    # Load existing students for duplicate checking
    existing_students = await db.db.students.find(
        {"institution_id": institution_id},
        {"student_id": 1, "roll_number": 1, "emails": 1, "contact": 1}
    ).to_list(length=10000)

    existing_rolls = set()
    existing_inst_emails = set()
    existing_personal_emails = set()

    for st in existing_students:
        r = st.get("student_id") or st.get("roll_number")
        if r: existing_rolls.add(str(r).strip().lower())
        
        ie = st.get("contact", {}).get("institutional_email") or st.get("emails", {}).get("institute")
        if ie: existing_inst_emails.add(str(ie).strip().lower())
        
        pe = st.get("contact", {}).get("personal_email") or st.get("emails", {}).get("personal")
        if pe: existing_personal_emails.add(str(pe).strip().lower())

    seen_batch_rolls = set()
    seen_batch_emails = set()

    batch_id = str(uuid.uuid4())
    canonical_records: List[CanonicalStudent] = []
    validation_issues: List[Dict[str, Any]] = []

    total_rows = 0
    valid_rows = 0
    duplicate_rows = 0
    missing_email_count = 0
    invalid_phone_count = 0
    unknown_dept_count = 0
    unknown_class_count = 0
    invalid_academic_score_count = 0
    requires_manual_review_count = 0
    batch_mismatch_count = 0

    header_indices = {h.lower(): i for i, h in enumerate(header_row) if h}

    def get_val(row: Tuple, *col_names: str) -> str:
        # 1. Exact match
        for name in col_names:
            idx = header_indices.get(name.lower().strip())
            if idx is not None and idx < len(row):
                v = clean_cell_str(row[idx])
                if v and v not in ["-", "--", "NA", "N/A", "null", "none", "nat"]:
                    return v
        # 2. Substring / case-insensitive search
        for name in col_names:
            norm_target = name.lower().strip()
            for h_lower, idx in header_indices.items():
                if (norm_target in h_lower or h_lower in norm_target) and idx < len(row):
                    v = clean_cell_str(row[idx])
                    if v and v not in ["-", "--", "NA", "N/A", "null", "none", "nat"]:
                        return v
        return ""

    now_utc = datetime.now(timezone.utc)

    for row_idx, row in enumerate(rows, start=2):
        if not any(row):
            continue

        total_rows += 1
        row_issues = []

        # 1. Roll Number
        roll_no = get_val(row, "Student ID / Roll Number", "Roll Number", "Register Number", "Roll No")
        if not roll_no:
            row_issues.append("Missing Student ID / Roll Number")

        # 2. Names
        full_name = get_val(row, "Full Name", "Name of Student ( As per 10th certificate)", "Name", "Student Name")
        first_name = get_val(row, "First Name")
        middle_name = get_val(row, "Middle Name")
        last_name = get_val(row, "Last Name")
        if not full_name and first_name:
            full_name = f"{first_name} {middle_name} {last_name}".strip()
        if not full_name:
            full_name = roll_no or f"Student-{row_idx}"
            row_issues.append("Missing Student Full Name")

        # 3. Emails & Mobile
        inst_email = clean_email(get_val(row, "Institutional Email", "Domain mail id", "Domain Email", "Email Address", "Official Email"))
        personal_email = clean_email(get_val(row, "Personal Email", "E mail ID ( Personal )", "Alternate Email", "Personal Mail"))
        mobile = clean_phone(get_val(row, "Mobile Number", "Student Mobile Number", "Mobile", "Contact Number"))

        if not inst_email and not personal_email:
            missing_email_count += 1
            row_issues.append("Missing both Institutional and Personal email")

        if mobile and len(mobile) < 10:
            invalid_phone_count += 1
            row_issues.append("Invalid mobile number format")

        # Duplicate detection
        is_duplicate = False
        if roll_no and roll_no.lower() in existing_rolls:
            is_duplicate = True
            row_issues.append(f"Roll number '{roll_no}' already exists in database")
        elif roll_no and roll_no.lower() in seen_batch_rolls:
            is_duplicate = True
            row_issues.append(f"Duplicate roll number '{roll_no}' inside upload sheet")

        if inst_email and inst_email.lower() in existing_inst_emails:
            is_duplicate = True
            row_issues.append(f"Institutional email '{inst_email}' already registered")
        elif inst_email and inst_email.lower() in seen_batch_emails:
            is_duplicate = True
            row_issues.append(f"Duplicate email '{inst_email}' inside upload sheet")

        if roll_no: seen_batch_rolls.add(roll_no.lower())
        if inst_email: seen_batch_emails.add(inst_email.lower())

        if is_duplicate:
            duplicate_rows += 1

        # 4. Academic Batch & Hierarchy Resolution
        prog_name = get_val(row, "Program", "Course", "Degree")
        dept_name = get_val(row, "Department", "Branch", "Stream")
        class_name = get_val(row, "Class / Section", "Class", "Section", "Branch & Class")
        
        # Batch & passing years
        grad_year_str = get_val(row, "Expected Graduation Year", "Year of Passing", "Graduation Year", "Passout Year")
        adm_year_str = get_val(row, "Admission Year")
        batch_label_str = get_val(row, "Academic Batch")

        hierarchy_res = await resolve_hierarchy_batch_and_faculty(
            institution_id, prog_name, dept_name, class_name,
            adm_year_str, grad_year_str, batch_label_str,
            all_programs, all_depts, all_classes, all_faculty
        )

        if hierarchy_res["batch_duration_mismatch"]:
            batch_mismatch_count += 1
            row_issues.append(f"BATCH_DURATION_MISMATCH: Admission ({hierarchy_res['admission_year']}) to Graduation ({hierarchy_res['expected_graduation_year']}) does not match {hierarchy_res['duration_years']}-year program duration")

        if not hierarchy_res["department_id"]:
            unknown_dept_count += 1
            row_issues.append(f"Could not resolve department '{dept_name}'")

        if not hierarchy_res["class_id"]:
            unknown_class_count += 1
            row_issues.append(f"Could not resolve class '{class_name}'")

        # 5. Demographics
        dob = get_val(row, "Date of Birth", "DOB")
        gender = get_val(row, "Gender", "Gender 2") or "Not Specified"
        nationality = get_val(row, "Nationality") or "Indian"

        # 6. Academic Score Normalization (10th, 12th/Diploma, UG Overall & Semesters)
        t_board = get_val(row, "10th Board", "Xth BOARD")
        t_inst = get_val(row, "10th School / Institution", "Xth School Name")
        t_year = get_val(row, "10th Passing Year", "Xth Year of passing")
        t_type = get_val(row, "10th Score Type", "Xth (% or CGPA) [Row 1]")
        t_score_raw = get_val(row, "10th Score", "Xth (% or CGPA)", "Xth MARKS")
        t_norm = normalize_academic_score(t_score_raw, expected_type=t_type)
        t_rule = await find_matching_conversion_rule(t_board, "10th", t_year, t_norm.get("score_type"))
        t_conv = apply_conversion_rule(t_rule, t_score_raw, t_norm.get("normalized_cgpa"), t_norm.get("normalized_percentage"))

        tw_board = get_val(row, "12th / Diploma Board", "XIIth / Diploma Board", "XIIth Board")
        tw_inst = get_val(row, "12th / Diploma Institution", "XIIth/ Diploma Institute Name", "XIIth Institute")
        tw_year = get_val(row, "12th / Diploma Passing Year", "XIIth year of passing")
        tw_type = get_val(row, "12th / Diploma Score Type", "XIIth or DIPLOMA (% or CGPA) [1.]")
        tw_score_raw = get_val(row, "12th / Diploma Score", "XIIth or DIPLOMA(% or CGPA)", "XIIth MARKS")
        tw_norm = normalize_academic_score(tw_score_raw, expected_type=tw_type)
        tw_rule = await find_matching_conversion_rule(tw_board, "12th", tw_year, tw_norm.get("score_type"))
        tw_conv = apply_conversion_rule(tw_rule, tw_score_raw, tw_norm.get("normalized_cgpa"), tw_norm.get("normalized_percentage"))

        ug_inst = get_val(row, "UG College / Institution", "College Name")
        ug_score_type = get_val(row, "UG Overall Score Type", "UG Total CGPA", "U.G. (% or CGPA)")
        ug_score_raw = get_val(row, "UG Overall Score", "UG Total CGPA", "U.G. (% or CGPA), if not there put -")
        ug_norm = normalize_academic_score(ug_score_raw, expected_type="CGPA")

        # Semesters SGPA mapping
        sgpa_dict: Dict[str, Optional[float]] = {}
        raw_sgpas: Dict[str, Optional[str]] = {}
        for sem_num, sem_key in [(1, "1-1"), (2, "1-2"), (3, "2-1"), (4, "2-2"), (5, "3-1"), (6, "3-2"), (7, "4-1"), (8, "4-2")]:
            raw_s = get_val(row, f"Semester {sem_num} SGPA", f"I-{sem_num} SGPA, if not there put -", f"II-{sem_num-2} SGPA, if not there put -", f"III-{sem_num-4} SGPA, if not there put -", f"IV-{sem_num-6} SGPA, if not there put -")
            raw_sgpas[sem_key] = raw_s or None
            s_norm = normalize_academic_score(raw_s, expected_type="CGPA")
            sgpa_dict[sem_key] = s_norm.get("normalized_cgpa")

        # Flags & Eligibility
        backlog_count_raw = get_val(row, "Backlog Count", "Current Backlogs", "No. of Backlogs")
        curr_backlogs = 0
        try:
            if backlog_count_raw: curr_backlogs = int(float(backlog_count_raw))
        except ValueError:
            curr_backlogs = 0

        edu_gap = get_val(row, "Education Gap", "Whether any Gaps in Education (Y/N)") or "No"
        first_attempt = get_val(row, "First Attempt Status", "Whether cleared all papers in first attempt or not (Y/N) ") or "Yes"

        # Check if academic values require review
        requires_review = (
            t_norm.get("verification_status") == "REQUIRES_REVIEW" or
            tw_norm.get("verification_status") == "REQUIRES_REVIEW" or
            ug_norm.get("verification_status") == "REQUIRES_REVIEW" or
            not hierarchy_res["is_mapped"] or
            hierarchy_res["batch_duration_mismatch"]
        )

        if requires_review:
            requires_manual_review_count += 1
            if t_norm.get("verification_status") == "REQUIRES_REVIEW":
                invalid_academic_score_count += 1
                row_issues.append(f"10th score '{t_score_raw}' requires manual review")
            if tw_norm.get("verification_status") == "REQUIRES_REVIEW":
                invalid_academic_score_count += 1
                row_issues.append(f"12th score '{tw_score_raw}' requires manual review")
            if ug_norm.get("verification_status") == "REQUIRES_REVIEW":
                invalid_academic_score_count += 1
                row_issues.append(f"UG score '{ug_score_raw}' requires manual review")

        if not row_issues:
            valid_rows += 1

        # 7. Compute System-Derived Study Year and Semester
        temp_student_dict = {
            "institution_id": institution_id,
            "admission_year": hierarchy_res["admission_year"],
            "expected_graduation_year": hierarchy_res["expected_graduation_year"],
            "academic_identity": {
                "program_id": hierarchy_res["program_id"],
                "admission_year": hierarchy_res["admission_year"],
                "expected_graduation_year": hierarchy_res["expected_graduation_year"]
            }
        }
        derived_status = await get_student_academic_status(temp_student_dict, current_date=now_utc, institute_id=institution_id)

        # 8. Build Canonical Model Object
        student_obj = CanonicalStudent(
            institute_id=institution_id,
            student_id=roll_no or f"TEMP-{row_idx}",
            identity=IdentityData(
                full_name=full_name,
                first_name=first_name or None,
                middle_name=middle_name or None,
                last_name=last_name or None,
                date_of_birth=dob or None,
                gender=gender,
                nationality=nationality
            ),
            contact=ContactData(
                institutional_email=inst_email or f"student{row_idx}@institution.edu",
                personal_email=personal_email or None,
                mobile=mobile or None
            ),
            academic_identity=AcademicIdentityData(
                program_id=hierarchy_res["program_id"],
                department_id=hierarchy_res["department_id"],
                class_id=hierarchy_res["class_id"],
                program_name=hierarchy_res["program_name"],
                department_name=hierarchy_res["department_name"],
                class_name=hierarchy_res["class_name"],
                academic_batch_id=hierarchy_res["academic_batch_id"],
                batch_label=hierarchy_res["batch_label"],
                admission_year=hierarchy_res["admission_year"],
                expected_graduation_year=hierarchy_res["expected_graduation_year"],
                graduation_year=hierarchy_res["expected_graduation_year"],
                current_study_year=derived_status["current_study_year"],
                current_semester=derived_status["current_semester"]
            ),
            education=EducationData(
                secondary=ScoreRecord(
                    board=t_board or None,
                    institution=t_inst or None,
                    passing_year=t_year or None,
                    raw_score=t_score_raw or None,
                    score_type=t_norm.get("score_type"),
                    normalized_percentage=t_conv.get("normalized_percentage"),
                    normalized_cgpa=t_conv.get("normalized_cgpa"),
                    normalized_value=t_conv.get("normalized_percentage") or t_conv.get("normalized_cgpa"),
                    coverage=t_norm.get("coverage"),
                    normalization_method=t_norm.get("normalization_method"),
                    conversion_rule_id=t_conv.get("conversion_rule_id"),
                    verification_status=t_conv.get("verification_status", "VERIFIED")
                ) if (t_board or t_score_raw) else None,
                higher_secondary_or_diploma=ScoreRecord(
                    board=tw_board or None,
                    institution=tw_inst or None,
                    passing_year=tw_year or None,
                    raw_score=tw_score_raw or None,
                    score_type=tw_norm.get("score_type"),
                    normalized_percentage=tw_conv.get("normalized_percentage"),
                    normalized_cgpa=tw_conv.get("normalized_cgpa"),
                    normalized_value=tw_conv.get("normalized_percentage") or tw_conv.get("normalized_cgpa"),
                    coverage=tw_norm.get("coverage"),
                    normalization_method=tw_norm.get("normalization_method"),
                    conversion_rule_id=tw_conv.get("conversion_rule_id"),
                    verification_status=tw_conv.get("verification_status", "VERIFIED")
                ) if (tw_board or tw_score_raw) else None,
                undergraduate=UndergraduateEducation(
                    institution=ug_inst or None,
                    overall_raw_score=ug_score_raw or None,
                    overall_score_type=ug_score_type or None,
                    normalized_cgpa=ug_norm.get("normalized_cgpa") or 0.0,
                    normalized_percentage=ug_norm.get("normalized_percentage"),
                    coverage=ug_norm.get("coverage"),
                    semester_sgpa=sgpa_dict,
                    raw_sgpas=raw_sgpas,
                    verification_status=ug_norm.get("verification_status", "VERIFIED")
                )
            ),
            eligibility_data=EligibilityData(
                current_backlogs=curr_backlogs,
                backlog_count=curr_backlogs,
                education_gap="Yes" if edu_gap.lower().startswith("y") else "No",
                first_attempt_status="Yes" if first_attempt.lower().startswith("y") else "No"
            ),
            entrance_exams=[
                EntranceExamRecord(
                    exam_name=get_val(row, "Entrance Exam", "EAMCET/ECET/ICET") or "EAMCET",
                    rank=get_val(row, "Entrance Rank", "Rank", "EAMCET Rank (If not there put -)") or None
                )
            ] if get_val(row, "Entrance Exam", "EAMCET/ECET/ICET", "Rank") else [],
            skills=[s.strip() for s in get_val(row, "Softwares learnt", "Skills").split(",") if s.strip() and s.strip().lower() != "none"],
            documents=DocumentsData(
                resume=(
                    (get_val(row, "Combined Document URL", "Upload your Document in single PDF", "Resume URL", "Resume Link", "Resume", "CV Link", "CV URL", "Single PDF") or "").strip()
                    if any(prot in (get_val(row, "Combined Document URL", "Upload your Document in single PDF", "Resume URL", "Resume Link", "Resume", "CV Link", "CV URL", "Single PDF") or "") for prot in ["http://", "https://", "drive.google.com", "docs.google.com"])
                    else None
                ),
                profile_photo=(
                    (get_val(row, "Profile Photo URL", "upload your photo", "Photo URL", "Photo Link", "Passport Photograph", "Photograph") or "").strip()
                    if any(prot in (get_val(row, "Profile Photo URL", "upload your photo", "Photo URL", "Photo Link", "Passport Photograph", "Photograph") or "") for prot in ["http://", "https://", "drive.google.com", "docs.google.com"])
                    else None
                ),
                combined_documents=(
                    (get_val(row, "Combined Document URL", "Upload your Document in single PDF", "All Documents") or "").strip()
                    if any(prot in (get_val(row, "Combined Document URL", "Upload your Document in single PDF", "All Documents") or "") for prot in ["http://", "https://", "drive.google.com", "docs.google.com"])
                    else None
                ),
                technical_certificates=(
                    (get_val(row, "Technical Certificates URL", "Make a single file of all technical certificates", "Technical Certificates URL", "Technical Certificates", "Technical Certs") or "").strip()
                    if any(prot in (get_val(row, "Technical Certificates URL", "Make a single file of all technical certificates", "Technical Certificates URL", "Technical Certificates", "Technical Certs") or "") for prot in ["http://", "https://", "drive.google.com", "docs.google.com"])
                    else None
                ),
                achievement_certificates=(
                    (get_val(row, "Achievement Certificates URL", "Any other achievements certificates", "Achievement Certificates", "Achievements") or "").strip()
                    if any(prot in (get_val(row, "Achievement Certificates URL", "Any other achievements certificates", "Achievement Certificates", "Achievements") or "") for prot in ["http://", "https://", "drive.google.com", "docs.google.com"])
                    else None
                )
            ),
            authentication=AuthenticationState(
                account_status="IMPORTED_PENDING_FACULTY",
                first_login_completed=False,
                password_set=False,
                verified_institutional_email=False,
                verified_personal_email=False
            ),
            approval=ApprovalState(
                faculty_status="PENDING" if (hierarchy_res["is_mapped"] and not hierarchy_res["batch_duration_mismatch"]) else "MAPPING_REVIEW",
                faculty_id=hierarchy_res["faculty_id"],
                faculty_name=hierarchy_res["faculty_name"]
            ),
            import_metadata=ImportMetadata(
                source_file=source_filename,
                import_batch_id=batch_id,
                imported_at=datetime.now(timezone.utc).isoformat(),
                raw_row_reference={header_row[i]: str(row[i]) for i in range(min(len(header_row), len(row))) if row[i] is not None}
            )
        )

        canonical_records.append(student_obj)

        if row_issues:
            validation_issues.append({
                "row_number": row_idx,
                "roll_number": roll_no,
                "name": full_name,
                "issues": row_issues
            })

    # Save to database if confirm_save is True
    inserted_count = 0
    updated_count = 0

    if confirm_save and canonical_records:
        for student in canonical_records:
            mongo_doc = student.to_mongo_dict()
            
            filter_query = {
                "institution_id": institution_id,
                "$or": [
                    {"student_id": student.student_id},
                    {"roll_number": student.student_id},
                    {"emails.institute": student.contact.institutional_email}
                ]
            }
            res = await db.db.students.update_one(filter_query, {"$set": mongo_doc}, upsert=True)
            if res.upserted_id:
                inserted_count += 1
            else:
                updated_count += 1

        # Centralized audit log
        await db.db.audit_logs.insert_one({
            "action": "standard_student_excel_import",
            "institution_id": institution_id,
            "user_id": imported_by_user_id,
            "source_file": source_filename,
            "batch_id": batch_id,
            "total_rows": total_rows,
            "inserted": inserted_count,
            "updated": updated_count,
            "timestamp": datetime.now(timezone.utc).isoformat(),
            "status": "success"
        })

    return {
        "success": True,
        "batch_id": batch_id,
        "summary": {
            "total_rows": total_rows,
            "valid_rows": valid_rows,
            "duplicate_rows": duplicate_rows,
            "missing_email_count": missing_email_count,
            "invalid_phone_count": invalid_phone_count,
            "unknown_department_count": unknown_dept_count,
            "unknown_class_count": unknown_class_count,
            "batch_mismatch_count": batch_mismatch_count,
            "invalid_academic_score_count": invalid_academic_score_count,
            "requires_manual_review_count": requires_manual_review_count,
            "inserted_count": inserted_count,
            "updated_count": updated_count
        },
        "sample_preview": [
            {
                "student_id": s.student_id,
                "full_name": s.identity.full_name,
                "institutional_email": s.contact.institutional_email,
                "program": s.academic_identity.program_name,
                "academic_batch": s.academic_identity.batch_label,
                "study_year": f"{s.academic_identity.current_study_year} Year",
                "semester": f"{s.academic_identity.current_semester} Sem",
                "department": s.academic_identity.department_name,
                "class_section": s.academic_identity.class_name,
                "assigned_faculty": s.approval.faculty_name or "Pending Assignment",
                "ug_cgpa": s.education.undergraduate.normalized_cgpa if s.education.undergraduate else None,
                "tenth_percentage": s.education.secondary.normalized_percentage if s.education.secondary else None,
                "twelfth_percentage": s.education.higher_secondary_or_diploma.normalized_percentage if s.education.higher_secondary_or_diploma else None,
                "status": s.approval.faculty_status
            }
            for s in canonical_records[:15]
        ],
        "validation_issues": validation_issues[:50]
    }

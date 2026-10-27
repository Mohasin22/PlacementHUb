from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File
from fastapi.responses import StreamingResponse
from app.db.tenant import get_current_user, get_tenant_id, RoleChecker
from app.db.connection import db
from pydantic import BaseModel
from typing import List, Optional, Dict, Any
from datetime import datetime, timezone
from app.db.utils import id_query, parse_id
from app.services.audit import log_audit_event
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
import io

router = APIRouter()

VALID_STATUSES = ["Applied", "Shortlisted", "Selected", "Cancelled", "Placed", "Rejected", "On-Hold"]

class ApplicationStatusUpdate(BaseModel):
    status: str
    reason: Optional[str] = None

class CancelSelectionPayload(BaseModel):
    reason: Optional[str] = "Selection cancelled by TPO"

class ExportFieldsPayload(BaseModel):
    fields: Optional[List[str]] = None
    status_filter: Optional[str] = None
    selected_application_ids: Optional[List[str]] = None

class BulkStatusUpdate(BaseModel):
    application_ids: List[str]
    status: str

from app.services.eligibility_engine import evaluate_student_eligibility_detailed

class PlaceStudentPayload(BaseModel):
    offer_letter_url: Optional[str] = None
    ctc_offered: Optional[str] = None
    joining_date: Optional[str] = None

class SelectStudentsForDrivePayload(BaseModel):
    student_ids: List[str]
    status: Optional[str] = "Selected"

# ── TPO Hierarchy Scoped to Program ──────────────────────────────────────────
@router.get("/hierarchy", dependencies=[Depends(RoleChecker(["TPO"]))])
async def get_tpo_hierarchy(
    current_user: dict = Depends(get_current_user),
    institution_id: str = Depends(get_tenant_id),
):
    """Fetch hierarchy scoped specifically to the TPO's assigned program."""
    tpo_user = await db.db.users.find_one(id_query(current_user["user_id"]))
    if not tpo_user or not tpo_user.get("program_id"):
        # Fallback to institution full hierarchy
        progs_cursor = db.db.programs.find({"$or": [{"institution_id": institution_id}, {"institution_id": parse_id(institution_id)}]})
        programs = await progs_cursor.to_list(length=50)
        hierarchy = []
        for prog in programs:
            p_id = str(prog.get("id") or prog["_id"])
            depts_cursor = db.db.departments.find({"$or": [{"program_id": p_id}, {"program_id": parse_id(p_id)}]})
            depts = await depts_cursor.to_list(length=100)
            departments_data = []
            for dept in depts:
                d_id = str(dept.get("id") or dept["_id"])
                classes_cursor = db.db.classes.find({"$or": [{"department_id": d_id}, {"department_id": parse_id(d_id)}]})
                classes = await classes_cursor.to_list(length=100)
                departments_data.append({
                    "id": d_id,
                    "name": dept.get("department_name") or dept.get("name", "Unknown"),
                    "classes": [{"id": str(c.get("id") or c["_id"]), "name": c.get("class_name") or c.get("name", "Unknown")} for c in classes]
                })
            hierarchy.append({
                "id": p_id,
                "name": prog.get("program_name") or prog.get("name", "Unknown"),
                "departments": departments_data
            })
        return {"program_id": None, "program_name": "All Programs", "hierarchy": hierarchy}
        
    prog_id = str(tpo_user["program_id"])
    prog = await db.db.programs.find_one(id_query(prog_id))
    prog_name = prog.get("program_name") or prog.get("name", "Unknown Program") if prog else "Program"
    
    depts_cursor = db.db.departments.find({
        "institution_id": institution_id,
        "$or": [{"program_id": prog_id}, {"program_id": parse_id(prog_id)}]
    })
    depts = await depts_cursor.to_list(length=100)
    
    departments_data = []
    for dept in depts:
        dept_id = str(dept.get("id") or dept["_id"])
        classes_cursor = db.db.classes.find({
            "$or": [{"department_id": dept_id}, {"department_id": parse_id(dept_id)}]
        })
        classes = await classes_cursor.to_list(length=100)
        
        departments_data.append({
            "id": dept_id,
            "name": dept.get("department_name") or dept.get("name", "Unknown Department"),
            "classes": [{"id": str(c.get("id") or c["_id"]), "name": c.get("class_name") or c.get("name", "Unknown Class")} for c in classes]
        })
        
    hierarchy = [{
        "id": prog_id,
        "name": prog_name,
        "departments": departments_data
    }]
    
    return {
        "program_id": prog_id,
        "program_name": prog_name,
        "hierarchy": hierarchy
    }

# ── TPO Drive Listing ────────────────────────────────────────────────────────
@router.get("/drives", dependencies=[Depends(RoleChecker(["TPO"]))])
async def list_tpo_drives(
    current_user: dict = Depends(get_current_user),
    institution_id: str = Depends(get_tenant_id),
):
    query = {"institution_id": institution_id}
    drives_cursor = db.db.placement_drives.find(query).sort("created_at", -1)
    drives = await drives_cursor.to_list(length=200)

    # 1. Aggregate application counts for all drives in one single indexed query
    pipeline = [
        {"$match": {"institution_id": institution_id}},
        {
            "$group": {
                "_id": {
                    "drive_id": "$drive_id",
                    "status": "$status"
                },
                "count": {"$sum": 1}
            }
        }
    ]
    app_stats = await db.db.applications.aggregate(pipeline).to_list(length=1000)
    stats_map: Dict[str, Dict[str, int]] = {}
    for item in app_stats:
        d_id = str(item["_id"].get("drive_id") or "")
        st = str(item["_id"].get("status") or "")
        cnt = int(item.get("count", 0))
        if d_id not in stats_map:
            stats_map[d_id] = {"total": 0, "Shortlisted": 0, "Placed": 0, "Rejected": 0}
        stats_map[d_id]["total"] += cnt
        if st in stats_map[d_id]:
            stats_map[d_id][st] += cnt

    now = datetime.now(timezone.utc)
    result = []
    for d in drives:
        drive_id = str(d["_id"])
        counts = stats_map.get(drive_id, {"total": 0, "Shortlisted": 0, "Placed": 0, "Rejected": 0})

        deadline = d.get("drive_deadline")
        parsed_deadline = None
        if deadline is not None:
            if isinstance(deadline, datetime):
                parsed_deadline = deadline if deadline.tzinfo is not None else deadline.replace(tzinfo=timezone.utc)
            elif isinstance(deadline, str):
                try:
                    clean_str = deadline.strip().replace("Z", "+00:00")
                    dt = datetime.fromisoformat(clean_str)
                    parsed_deadline = dt if dt.tzinfo is not None else dt.replace(tzinfo=timezone.utc)
                except ValueError:
                    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%d-%m-%Y"):
                        try:
                            parsed_deadline = datetime.strptime(deadline.strip(), fmt).replace(tzinfo=timezone.utc)
                            break
                        except ValueError:
                            pass

        is_active = (parsed_deadline > now) if parsed_deadline else True

        # Provide a safe string representation for the frontend
        dl_display = None
        if parsed_deadline:
            dl_display = parsed_deadline.isoformat()
        elif deadline:
            dl_display = str(deadline)

        result.append({
            "id": drive_id,
            "company_name": d.get("company_name"),
            "job_role": d.get("job_role"),
            "package": d.get("package"),
            "location": d.get("location"),
            "mode": d.get("mode"),
            "min_cgpa": d.get("min_cgpa"),
            "max_backlogs": d.get("max_backlogs"),
            "drive_deadline": dl_display,
            "is_active": is_active,
            "total_applicants": counts["total"],
            "shortlisted": counts["Shortlisted"],
            "placed": counts["Placed"],
            "rejected": counts["Rejected"],
            "external_apply_link": d.get("external_apply_link"),
            "gender_filter": d.get("gender_filter", "All"),
            "passout_year": d.get("passout_year", datetime.now().year),
            "allowed_departments": d.get("allowed_departments", []),
            "allowed_programs": d.get("allowed_programs", [])
        })
    return result

from app.services.academic_timeline_service import get_student_academic_status

# ── TPO Students Listing ─────────────────────────────────────────────────────
@router.get("/students", dependencies=[Depends(RoleChecker(["TPO", "Dean"]))])
async def get_tpo_students(
    program_id: Optional[str] = None,
    department_id: Optional[str] = None,
    academic_batch_id: Optional[str] = None,
    study_year: Optional[int] = None,
    semester: Optional[int] = None,
    class_id: Optional[str] = None,
    search: Optional[str] = None,
    current_user: dict = Depends(get_current_user),
    institution_id: str = Depends(get_tenant_id),
):
    query: Dict[str, Any] = {"institution_id": institution_id}
    if program_id:
        query["$or"] = [{"program_id": program_id}, {"academic_identity.program_id": program_id}]
    if department_id:
        query["$and"] = query.get("$and", []) + [{"$or": [{"department_id": department_id}, {"academic_identity.department_id": department_id}]}]
    if academic_batch_id:
        query["$and"] = query.get("$and", []) + [{"$or": [{"academic_batch_id": academic_batch_id}, {"academic_identity.academic_batch_id": academic_batch_id}]}]
    if class_id:
        query["$and"] = query.get("$and", []) + [{"$or": [{"class_id": class_id}, {"academic_identity.class_id": class_id}]}]

    students_cursor = db.db.students.find(query).sort("student_id", 1)
    students = await students_cursor.to_list(length=2000)
    
    now_utc = datetime.now(timezone.utc)
    result = []
    for st in students:
        # Calculate derived status
        derived = await get_student_academic_status(st, current_date=now_utc, institute_id=institution_id)
        
        # Filter by study year if requested
        if study_year is not None and derived["current_study_year"] != study_year:
            continue
        # Filter by semester if requested
        if semester is not None and derived["current_semester"] != semester:
            continue

        name = st.get("name") or st.get("identity", {}).get("full_name") or ""
        roll = st.get("roll_number") or st.get("student_id") or ""
        
        if search:
            s_low = search.lower().strip()
            if s_low not in name.lower() and s_low not in roll.lower():
                continue

        st_data = {
            "id": str(st["_id"]),
            "name": name,
            "roll_number": roll,
            "emails": {
                "institute": st.get("contact", {}).get("institutional_email") or st.get("emails", {}).get("institute") or "",
                "personal": st.get("contact", {}).get("personal_email") or st.get("emails", {}).get("personal") or ""
            },
            "mobile": st.get("mobile") or st.get("contact", {}).get("mobile") or "",
            "status": st.get("status") or st.get("approval", {}).get("faculty_status") or "APPROVED",
            "cgpa": st.get("cgpa") or st.get("education", {}).get("undergraduate", {}).get("normalized_cgpa") or 0.0,
            "active_backlogs": st.get("active_backlogs") if st.get("active_backlogs") is not None else st.get("eligibility_data", {}).get("current_backlogs", 0),
            "total_backlogs": st.get("total_backlogs") if st.get("total_backlogs") is not None else st.get("eligibility_data", {}).get("backlog_count", 0),
            "program_id": str(st.get("program_id") or st.get("academic_identity", {}).get("program_id") or ""),
            "program_name": st.get("program_name") or st.get("academic_identity", {}).get("program_name") or "",
            "department_id": str(st.get("department_id") or st.get("academic_identity", {}).get("department_id") or ""),
            "department_name": st.get("department_name") or st.get("academic_identity", {}).get("department_name") or "General",
            "class_id": str(st.get("class_id") or st.get("academic_identity", {}).get("class_id") or ""),
            "class_name": st.get("class_name") or st.get("academic_identity", {}).get("class_name") or "",
            "academic_batch_id": st.get("academic_batch_id") or st.get("academic_identity", {}).get("academic_batch_id") or "",
            "batch_label": derived["batch_label"],
            "current_study_year": derived["current_study_year"],
            "current_semester": derived["current_semester"],
            "study_year_label": derived["study_year_label"],
            "semester_label": derived["semester_label"],
            "skills": st.get("skills", []),
            "projects": st.get("projects", []),
            "certifications": st.get("certifications", []),
            "internships": st.get("internships", []),
            "achievements": st.get("achievements", []),
            "entrance_exams": st.get("entrance_exams", []),
            "links": st.get("links", {}),
            "documents": st.get("documents", {}),
            "identity": st.get("identity", {}),
            "contact": st.get("contact", {}),
            "education": st.get("education", {}),
            "eligibility_data": st.get("eligibility_data", {}),
            "resume_url": st.get("resume_url") or st.get("documents", {}).get("resume") or "",
            "photo_url": st.get("photo_url") or st.get("documents", {}).get("profile_photo") or "",
            "dob": st.get("dob") or st.get("identity", {}).get("date_of_birth") or "",
            "gender": st.get("gender") or st.get("identity", {}).get("gender") or "Not Specified",
            "nationality": st.get("nationality") or st.get("identity", {}).get("nationality") or "Indian"
        }
        result.append(st_data)
        
    return result

@router.get("/students/{student_id}", dependencies=[Depends(RoleChecker(["TPO", "Dean"]))])
async def get_tpo_student_detail(
    student_id: str,
    current_user: dict = Depends(get_current_user),
    institution_id: str = Depends(get_tenant_id),
):
    """Retrieve full profile details of an individual student for TPO."""
    st = await db.db.students.find_one(id_query(student_id))
    if not st:
        raise HTTPException(status_code=404, detail="Student not found.")
        
    derived = await get_student_academic_status(st, current_date=datetime.now(timezone.utc), institute_id=institution_id)
    name = st.get("name") or st.get("identity", {}).get("full_name") or ""
    roll = st.get("roll_number") or st.get("student_id") or ""
    
    return {
        "id": str(st["_id"]),
        "name": name,
        "roll_number": roll,
        "emails": {
            "institute": st.get("contact", {}).get("institutional_email") or st.get("emails", {}).get("institute") or "",
            "personal": st.get("contact", {}).get("personal_email") or st.get("emails", {}).get("personal") or ""
        },
        "mobile": st.get("mobile") or st.get("contact", {}).get("mobile") or "",
        "status": st.get("status") or st.get("approval", {}).get("faculty_status") or "APPROVED",
        "cgpa": st.get("cgpa") or st.get("education", {}).get("undergraduate", {}).get("normalized_cgpa") or 0.0,
        "active_backlogs": st.get("active_backlogs") if st.get("active_backlogs") is not None else st.get("eligibility_data", {}).get("current_backlogs", 0),
        "total_backlogs": st.get("total_backlogs") if st.get("total_backlogs") is not None else st.get("eligibility_data", {}).get("backlog_count", 0),
        "program_id": str(st.get("program_id") or st.get("academic_identity", {}).get("program_id") or ""),
        "program_name": st.get("program_name") or st.get("academic_identity", {}).get("program_name") or "",
        "department_id": str(st.get("department_id") or st.get("academic_identity", {}).get("department_id") or ""),
        "department_name": st.get("department_name") or st.get("academic_identity", {}).get("department_name") or "General",
        "class_id": str(st.get("class_id") or st.get("academic_identity", {}).get("class_id") or ""),
        "class_name": st.get("class_name") or st.get("academic_identity", {}).get("class_name") or "",
        "academic_batch_id": st.get("academic_batch_id") or st.get("academic_identity", {}).get("academic_batch_id") or "",
        "batch_label": derived["batch_label"],
        "current_study_year": derived["current_study_year"],
        "current_semester": derived["current_semester"],
        "study_year_label": derived["study_year_label"],
        "semester_label": derived["semester_label"],
        "skills": st.get("skills", []),
        "projects": st.get("projects", []),
        "certifications": st.get("certifications", []),
        "internships": st.get("internships", []),
        "achievements": st.get("achievements", []),
        "entrance_exams": st.get("entrance_exams", []),
        "links": st.get("links", {}),
        "documents": st.get("documents", {}),
        "identity": st.get("identity", {}),
        "contact": st.get("contact", {}),
        "education": st.get("education", {}),
        "eligibility_data": st.get("eligibility_data", {}),
        "resume_url": st.get("resume_url") or st.get("documents", {}).get("resume") or "",
        "photo_url": st.get("photo_url") or st.get("documents", {}).get("profile_photo") or "",
        "dob": st.get("dob") or st.get("identity", {}).get("date_of_birth") or "",
        "gender": st.get("gender") or st.get("identity", {}).get("gender") or "Not Specified",
        "nationality": st.get("nationality") or st.get("identity", {}).get("nationality") or "Indian"
    }

# ── Drive Applications Listing with Snapshots ────────────────────────────────
@router.get("/drives/{drive_id}/applications", dependencies=[Depends(RoleChecker(["TPO"]))])
async def get_drive_applications(
    drive_id: str,
    current_user: dict = Depends(get_current_user),
    institution_id: str = Depends(get_tenant_id),
):
    drive = await db.db.placement_drives.find_one(id_query(drive_id))
    if not drive:
        raise HTTPException(status_code=404, detail="Placement drive not found.")
        
    apps_cursor = db.db.applications.find({"drive_id": drive_id, "institution_id": institution_id}).sort("applied_at", -1)
    apps = await apps_cursor.to_list(length=1000)
    
    app_details = []
    for app in apps:
        student = await db.db.students.find_one(id_query(app["student_id"]))
        # Use snapshot data if present, otherwise fall back to live profile
        p_snap = app.get("profile_snapshot") or {}
        a_snap = app.get("academic_snapshot") or {}
        
        name = p_snap.get("name") or (student.get("name") if student else "Unknown")
        roll_no = p_snap.get("roll_number") or (student.get("roll_number") if student else "")
        dept = p_snap.get("department") or (student.get("department_name") if student else "")
        cgpa = a_snap.get("ug_cgpa") or (student.get("cgpa") if student else 0.0)
        mobile = p_snap.get("mobile") or (student.get("mobile") if student else "")
        resume = app.get("resume_version") or p_snap.get("resume_url") or (student.get("resume_url") if student else "")
        skills = p_snap.get("skills") or (student.get("skills", []) if student else [])
        
        app_details.append({
            "application_id": str(app["_id"]),
            "student_id": app["student_id"],
            "name": name,
            "roll_number": roll_no,
            "department_name": dept,
            "cgpa": cgpa,
            "mobile": mobile,
            "resume_url": resume,
            "skills": skills,
            "applied_at": app.get("applied_at"),
            "status": app.get("status", "Applied"),
            "status_history": app.get("status_history", []),
            "offer_letter_url": app.get("offer_letter_url"),
            "ctc_offered": app.get("ctc_offered"),
            "profile_snapshot": p_snap,
            "academic_snapshot": a_snap
        })
        
    return app_details

# ── Application Status Updates & Selection Workflow ──────────────────────────

@router.post("/applications/{application_id}/status", dependencies=[Depends(RoleChecker(["TPO"]))])
async def update_application_status(
    application_id: str,
    payload: ApplicationStatusUpdate,
    current_user: dict = Depends(get_current_user),
    institution_id: str = Depends(get_tenant_id),
):
    if payload.status not in VALID_STATUSES:
        raise HTTPException(status_code=400, detail=f"Invalid status. Valid values: {VALID_STATUSES}")
    app = await db.db.applications.find_one(id_query(application_id))
    if not app:
        raise HTTPException(status_code=404, detail="Application record not found.")

    now_iso = datetime.now(timezone.utc).isoformat()
    history_entry = {
        "status": payload.status,
        "timestamp": now_iso,
        "changed_by": current_user["email"],
        "reason": payload.reason or f"Status transitioned to {payload.status}"
    }
        
    await db.db.applications.update_one(
        id_query(application_id),
        {
            "$set": {"status": payload.status, "updated_at": now_iso},
            "$push": {"status_history": history_entry}
        }
    )
    
    await log_audit_event(
        action=f"application_status_{payload.status.lower()}",
        actor=current_user["email"],
        actor_role="TPO",
        entity_type="Application",
        entity_id=application_id,
        institution_id=institution_id,
        old_value=app.get("status"),
        new_value=payload.status,
        details={"reason": payload.reason}
    )
    
    return {"message": f"Application status updated to {payload.status}."}

@router.post("/applications/{application_id}/select", dependencies=[Depends(RoleChecker(["TPO"]))])
async def select_application(
    application_id: str,
    current_user: dict = Depends(get_current_user),
    institution_id: str = Depends(get_tenant_id)
):
    """Marks student application as 'Selected' and records in status history."""
    app = await db.db.applications.find_one(id_query(application_id))
    if not app:
        raise HTTPException(status_code=404, detail="Application not found.")

    now_iso = datetime.now(timezone.utc).isoformat()
    history_entry = {
        "status": "Selected",
        "timestamp": now_iso,
        "changed_by": current_user["email"],
        "reason": "Selected for placement"
    }

    await db.db.applications.update_one(
        id_query(application_id),
        {
            "$set": {"status": "Selected", "updated_at": now_iso},
            "$push": {"status_history": history_entry}
        }
    )

    await log_audit_event(
        action="application_selected",
        actor=current_user["email"],
        actor_role="TPO",
        entity_type="Application",
        entity_id=application_id,
        institution_id=institution_id,
        old_value=app.get("status"),
        new_value="Selected"
    )

    return {"message": "Student application marked as Selected.", "status": "Selected"}

@router.post("/applications/{application_id}/cancel-selection", dependencies=[Depends(RoleChecker(["TPO"]))])
async def cancel_application_selection(
    application_id: str,
    payload: Optional[CancelSelectionPayload] = None,
    current_user: dict = Depends(get_current_user),
    institution_id: str = Depends(get_tenant_id)
):
    """
    Cancels student selection without deleting the record.
    Marks status as 'Cancelled' and records reason in status history.
    """
    app = await db.db.applications.find_one(id_query(application_id))
    if not app:
        raise HTTPException(status_code=404, detail="Application not found.")

    now_iso = datetime.now(timezone.utc).isoformat()
    reason = payload.reason if payload and payload.reason else "Selection cancelled by TPO"
    history_entry = {
        "status": "Cancelled",
        "timestamp": now_iso,
        "changed_by": current_user["email"],
        "reason": reason
    }

    await db.db.applications.update_one(
        id_query(application_id),
        {
            "$set": {"status": "Cancelled", "updated_at": now_iso},
            "$push": {"status_history": history_entry}
        }
    )

    await log_audit_event(
        action="application_selection_cancelled",
        actor=current_user["email"],
        actor_role="TPO",
        entity_type="Application",
        entity_id=application_id,
        institution_id=institution_id,
        old_value=app.get("status"),
        new_value="Cancelled",
        details={"reason": reason}
    )

    return {"message": "Student selection cancelled.", "status": "Cancelled"}

@router.post("/drives/{drive_id}/bulk-select", dependencies=[Depends(RoleChecker(["TPO"]))])
async def bulk_select_applications(
    drive_id: str,
    payload: BulkStatusUpdate,
    current_user: dict = Depends(get_current_user),
    institution_id: str = Depends(get_tenant_id)
):
    """Bulk marks multiple applications as 'Selected'."""
    now_iso = datetime.now(timezone.utc).isoformat()
    updated = 0
    history_entry = {
        "status": "Selected",
        "timestamp": now_iso,
        "changed_by": current_user["email"],
        "reason": "Bulk selection"
    }

    for app_id in payload.application_ids:
        res = await db.db.applications.update_one(
            id_query(app_id),
            {
                "$set": {"status": "Selected", "updated_at": now_iso},
                "$push": {"status_history": history_entry}
            }
        )
        if res.modified_count:
            updated += 1

    await log_audit_event(
        action="bulk_applications_selected",
        actor=current_user["email"],
        actor_role="TPO",
        entity_type="PlacementDrive",
        entity_id=drive_id,
        institution_id=institution_id,
        details={"count": updated}
    )

    return {"message": f"{updated} application(s) selected.", "updated_count": updated}

@router.get("/drives/{drive_id}/candidate-pool", dependencies=[Depends(RoleChecker(["TPO", "Dean"]))])
async def get_drive_candidate_pool(
    drive_id: str,
    current_user: dict = Depends(get_current_user),
    institution_id: str = Depends(get_tenant_id)
):
    """Fetches candidate students across branches/years/sections with real-time eligibility evaluation for this drive."""
    drive = await db.db.placement_drives.find_one(id_query(drive_id))
    if not drive:
        raise HTTPException(status_code=404, detail="Placement drive not found.")

    students_cursor = db.db.students.find({
        "institution_id": institution_id,
        "status": {"$nin": ["blocked", "rejected", "inactive", "BLOCKED", "REJECTED"]}
    }).sort("student_id", 1)
    all_students = await students_cursor.to_list(length=3000)

    apps_cursor = db.db.applications.find({"drive_id": drive_id, "institution_id": institution_id})
    all_apps = await apps_cursor.to_list(length=3000)
    apps_by_student_id = {str(a["student_id"]): a for a in all_apps}

    now_utc = datetime.now(timezone.utc)
    from app.services.academic_timeline_service import get_student_academic_status, ensure_default_academic_calendar
    cached_cal = await db.db.academic_calendars.find_one({"institution_id": institution_id, "active": True})
    if not cached_cal:
        cached_cal = await ensure_default_academic_calendar(institution_id)

    raw_progs = await db.db.programs.find({"institution_id": institution_id}).to_list(length=100)
    cached_progs = {str(p["_id"]): p for p in raw_progs}
    for p in raw_progs:
        if "id" in p:
            cached_progs[str(p["id"])] = p

    candidates = []
    for st in all_students:
        st_id = str(st["_id"])
        derived = await get_student_academic_status(
            st,
            current_date=now_utc,
            institute_id=institution_id,
            cached_calendar=cached_cal,
            cached_programs=cached_progs
        )
        is_el, reasons, breakdown = evaluate_student_eligibility_detailed(st, drive, derived_status=derived)
        existing_app = apps_by_student_id.get(st_id)

        name = st.get("name") or st.get("identity", {}).get("full_name") or "Student"
        roll = st.get("roll_number") or st.get("student_id") or ""
        dept_name = st.get("department_name") or st.get("academic_identity", {}).get("department_name") or "General"
        dept_id = str(st.get("department_id") or st.get("academic_identity", {}).get("department_id") or dept_name)
        class_name = st.get("class_name") or st.get("academic_identity", {}).get("class_name") or ""
        class_id = str(st.get("class_id") or st.get("academic_identity", {}).get("class_id") or "")

        cgpa = float(st.get("cgpa") or st.get("education", {}).get("undergraduate", {}).get("normalized_cgpa") or 0.0)
        active_backlogs = int(st.get("active_backlogs") if st.get("active_backlogs") is not None else st.get("eligibility_data", {}).get("current_backlogs", 0))
        tenth_pct = st.get("education", {}).get("secondary", {}).get("normalized_percentage")
        twelfth_pct = st.get("education", {}).get("higher_secondary_or_diploma", {}).get("normalized_percentage")

        candidates.append({
            "student_id": st_id,
            "name": name,
            "roll_number": roll,
            "email": st.get("contact", {}).get("institutional_email") or st.get("emails", {}).get("institute") or st.get("contact", {}).get("personal_email") or "",
            "mobile": st.get("mobile") or st.get("contact", {}).get("mobile") or "",
            "department_id": dept_id,
            "department_name": dept_name,
            "class_id": class_id,
            "class_name": class_name,
            "current_study_year": derived.get("current_study_year", 4),
            "study_year_label": derived.get("study_year_label", "4th Year"),
            "passout_year": st.get("passout_year") or st.get("academic_identity", {}).get("expected_graduation_year") or st.get("academic_identity", {}).get("graduation_year"),
            "cgpa": cgpa,
            "active_backlogs": active_backlogs,
            "total_backlogs": int(st.get("total_backlogs") if st.get("total_backlogs") is not None else st.get("eligibility_data", {}).get("backlog_count", 0)),
            "tenth_percentage": tenth_pct,
            "twelfth_percentage": twelfth_pct,
            "gender": st.get("gender") or st.get("identity", {}).get("gender") or "Not Specified",
            "resume_url": st.get("resume_url") or st.get("documents", {}).get("resume") or "",
            "is_eligible": is_el,
            "eligibility_reasons": reasons,
            "eligibility_breakdown": breakdown,
            "application_status": existing_app.get("status") if existing_app else "Not Applied",
            "application_id": str(existing_app["_id"]) if existing_app else None
        })

    drive_summary = {
        "id": drive_id,
        "company_name": drive.get("company_name"),
        "job_role": drive.get("job_role"),
        "package": drive.get("package"),
        "min_cgpa": drive.get("min_cgpa", 0.0),
        "max_backlogs": drive.get("max_backlogs", 0),
        "min_tenth_percentage": drive.get("min_tenth_percentage"),
        "min_twelfth_percentage": drive.get("min_twelfth_percentage"),
        "passout_year": drive.get("passout_year"),
        "allowed_departments": drive.get("allowed_departments", []),
        "allowed_programs": drive.get("allowed_programs", []),
        "gender_filter": drive.get("gender_filter", "All"),
        "mode": drive.get("mode", "On-campus")
    }

    return {
        "drive": drive_summary,
        "total_candidates": len(candidates),
        "candidates": candidates
    }

@router.post("/drives/{drive_id}/select-students", dependencies=[Depends(RoleChecker(["TPO"]))])
async def select_students_for_drive(
    drive_id: str,
    payload: SelectStudentsForDrivePayload,
    current_user: dict = Depends(get_current_user),
    institution_id: str = Depends(get_tenant_id)
):
    """Enrolls/selects candidate students for a drive directly with candidate snapshot and status history."""
    drive = await db.db.placement_drives.find_one(id_query(drive_id))
    if not drive:
        raise HTTPException(status_code=404, detail="Placement drive not found.")

    if not payload.student_ids:
        raise HTTPException(status_code=400, detail="No student IDs provided.")

    target_status = payload.status or "Selected"
    if target_status not in VALID_STATUSES:
        target_status = "Selected"

    now_iso = datetime.now(timezone.utc).isoformat()
    selected_count = 0
    notifications_to_insert = []

    for st_id in payload.student_ids:
        st = await db.db.students.find_one(id_query(st_id))
        if not st:
            continue

        existing_app = await db.db.applications.find_one({"student_id": st_id, "drive_id": drive_id})
        history_entry = {
            "status": target_status,
            "timestamp": now_iso,
            "changed_by": current_user["email"],
            "reason": f"Marked as {target_status} via Excel Student Selection tool"
        }

        if existing_app:
            await db.db.applications.update_one(
                id_query(str(existing_app["_id"])),
                {
                    "$set": {"status": target_status, "updated_at": now_iso},
                    "$push": {"status_history": history_entry}
                }
            )
            selected_count += 1
        else:
            profile_snapshot = {
                "name": st.get("name") or st.get("identity", {}).get("full_name"),
                "roll_number": st.get("roll_number") or st.get("student_id"),
                "institutional_email": st.get("contact", {}).get("institutional_email") or st.get("emails", {}).get("institute"),
                "personal_email": st.get("contact", {}).get("personal_email") or st.get("emails", {}).get("personal"),
                "mobile": st.get("mobile") or st.get("contact", {}).get("mobile"),
                "program": st.get("program_name") or st.get("academic_identity", {}).get("program_name"),
                "department": st.get("department_name") or st.get("academic_identity", {}).get("department_name"),
                "class_section": st.get("class_name") or st.get("academic_identity", {}).get("class_name"),
                "passout_year": st.get("passout_year") or st.get("academic_identity", {}).get("graduation_year"),
                "skills": st.get("skills", []),
                "projects": st.get("projects", []),
                "internships": st.get("internships", []),
                "certifications": st.get("certifications", []),
                "links": st.get("links", {}),
                "resume_url": st.get("resume_url") or st.get("documents", {}).get("resume")
            }
            academic_snapshot = {
                "tenth_percentage": st.get("education", {}).get("secondary", {}).get("normalized_percentage"),
                "tenth_board": st.get("education", {}).get("secondary", {}).get("board"),
                "twelfth_percentage": st.get("education", {}).get("higher_secondary_or_diploma", {}).get("normalized_percentage"),
                "twelfth_board": st.get("education", {}).get("higher_secondary_or_diploma", {}).get("board"),
                "ug_cgpa": st.get("cgpa") or st.get("education", {}).get("undergraduate", {}).get("normalized_cgpa"),
                "ug_percentage": st.get("education", {}).get("undergraduate", {}).get("normalized_percentage"),
                "current_backlogs": st.get("active_backlogs") if st.get("active_backlogs") is not None else st.get("eligibility_data", {}).get("current_backlogs"),
                "total_backlogs": st.get("total_backlogs") if st.get("total_backlogs") is not None else st.get("eligibility_data", {}).get("backlog_count"),
                "education_gap": st.get("eligibility_data", {}).get("education_gap") or "No",
                "first_attempt_status": st.get("eligibility_data", {}).get("first_attempt_status") or "Yes"
            }
            new_app = {
                "student_id": st_id,
                "drive_id": drive_id,
                "institution_id": institution_id,
                "profile_snapshot": profile_snapshot,
                "academic_snapshot": academic_snapshot,
                "resume_version": profile_snapshot.get("resume_url"),
                "status": target_status,
                "status_history": [history_entry],
                "applied_at": now_iso,
                "selected_by_tpo": True
            }
            await db.db.applications.insert_one(new_app)
            selected_count += 1

        notifications_to_insert.append({
            "student_id": st_id,
            "drive_id": drive_id,
            "institution_id": institution_id,
            "message": f"You have been marked as {target_status} for {drive.get('company_name')} ({drive.get('job_role')})!",
            "status": "unread",
            "created_at": now_iso
        })

    if notifications_to_insert:
        await db.db.notifications.insert_many(notifications_to_insert)

    await log_audit_event(
        action="students_selected_for_drive",
        actor=current_user["email"],
        actor_role="TPO",
        entity_type="PlacementDrive",
        entity_id=drive_id,
        institution_id=institution_id,
        details={"company": drive.get("company_name"), "selected_count": selected_count, "target_status": target_status}
    )

    return {
        "message": f"Successfully marked {selected_count} student(s) as {target_status} for {drive.get('company_name')}.",
        "selected_count": selected_count,
        "status": target_status
    }

@router.post("/drives/{drive_id}/bulk-cancel-selection", dependencies=[Depends(RoleChecker(["TPO"]))])
async def bulk_cancel_selection_applications(
    drive_id: str,
    payload: BulkStatusUpdate,
    current_user: dict = Depends(get_current_user),
    institution_id: str = Depends(get_tenant_id)
):
    """Bulk cancels selection for multiple applications without deleting records."""
    now_iso = datetime.now(timezone.utc).isoformat()
    updated = 0
    history_entry = {
        "status": "Cancelled",
        "timestamp": now_iso,
        "changed_by": current_user["email"],
        "reason": "Bulk selection cancellation"
    }

    for app_id in payload.application_ids:
        res = await db.db.applications.update_one(
            id_query(app_id),
            {
                "$set": {"status": "Cancelled", "updated_at": now_iso},
                "$push": {"status_history": history_entry}
            }
        )
        if res.modified_count:
            updated += 1

    await log_audit_event(
        action="bulk_applications_cancelled",
        actor=current_user["email"],
        actor_role="TPO",
        entity_type="PlacementDrive",
        entity_id=drive_id,
        institution_id=institution_id,
        details={"count": updated}
    )

    return {"message": f"{updated} selection(s) cancelled.", "updated_count": updated}

@router.post("/drives/{drive_id}/bulk-status", dependencies=[Depends(RoleChecker(["TPO"]))])
async def bulk_update_status(
    drive_id: str,
    payload: BulkStatusUpdate,
    current_user: dict = Depends(get_current_user),
    institution_id: str = Depends(get_tenant_id),
):
    if payload.status not in VALID_STATUSES:
        raise HTTPException(status_code=400, detail=f"Invalid status. Valid values: {VALID_STATUSES}")
    if not payload.application_ids:
        raise HTTPException(status_code=400, detail="No application IDs provided.")
        
    now_iso = datetime.now(timezone.utc).isoformat()
    history_entry = {
        "status": payload.status,
        "timestamp": now_iso,
        "changed_by": current_user["email"],
        "reason": f"Bulk status update to {payload.status}"
    }

    updated = 0
    for app_id in payload.application_ids:
        try:
            res = await db.db.applications.update_one(
                id_query(app_id),
                {
                    "$set": {"status": payload.status, "updated_at": now_iso},
                    "$push": {"status_history": history_entry}
                }
            )
            if res.modified_count:
                updated += 1
        except Exception:
            continue
            
    await log_audit_event(
        action=f"bulk_status_{payload.status.lower()}",
        actor=current_user["email"],
        actor_role="TPO",
        entity_type="PlacementDrive",
        entity_id=drive_id,
        institution_id=institution_id,
        details={"updated_count": updated, "status": payload.status}
    )
    
    return {"message": f"{updated} application(s) updated to {payload.status}."}

# ── Sanitized HR Excel Export (Snapshot-Based, Excludes Sensitive Data) ───────
@router.post("/drives/{drive_id}/export", dependencies=[Depends(RoleChecker(["TPO", "Dean"]))])
async def export_drive_applications_for_hr(
    drive_id: str,
    payload: ExportFieldsPayload,
    current_user: dict = Depends(get_current_user),
    institution_id: str = Depends(get_tenant_id),
):
    """
    Generates sanitized HR Excel export derived directly from application snapshot records.
    Strictly excludes sensitive fields (Aadhaar, religion, caste, eyesight, parent financial).
    """
    drive = await db.db.placement_drives.find_one(id_query(drive_id))
    if not drive:
        raise HTTPException(status_code=404, detail="Drive not found.")
        
    query: Dict[str, Any] = {"drive_id": drive_id, "institution_id": institution_id}
    if payload.status_filter and payload.status_filter != "All":
        query["status"] = payload.status_filter
    if payload.selected_application_ids:
        query["$or"] = [
            {"_id": {"$in": [parse_id(i) for i in payload.selected_application_ids]}},
            {"application_id": {"$in": payload.selected_application_ids}}
        ]
        
    apps_cursor = db.db.applications.find(query).sort("applied_at", 1)
    apps = await apps_cursor.to_list(length=5000)

    # Standard Canonical HR Export Headers
    DEFAULT_HR_HEADERS = [
        "Student ID / Roll Number", "Full Name", "Institutional Email", "Personal Email", "Mobile Number",
        "Program", "Department", "Class / Section", "Year of Passing",
        "10th Percentage", "12th / Diploma Percentage", "UG CGPA", "UG Percentage",
        "Current Backlogs", "Skills", "Projects", "Internships", "Certifications",
        "LinkedIn", "GitHub", "Portfolio", "Resume URL", "Application Status"
    ]

    export_columns = payload.fields if payload.fields else DEFAULT_HR_HEADERS

    # Filter out any sensitive field requested by mistake
    SENSITIVE_FORBIDDEN_TERMS = ["aadhaar", "caste", "socialstatus", "religion", "hostel", "eyesight", "reimbursement", "password", "income"]
    safe_columns = []
    for col in export_columns:
        col_clean = col.lower().replace(" ", "").replace("_", "").replace("-", "")
        if not any(st in col_clean for st in SENSITIVE_FORBIDDEN_TERMS):
            safe_columns.append(col)

    wb = Workbook()
    ws = wb.active
    ws.title = "HR Shortlist"

    # Styling
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=10)
    
    ws.append(safe_columns)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for app in apps:
        p_snap = app.get("profile_snapshot") or {}
        a_snap = app.get("academic_snapshot") or {}
        
        # Fallback to student live record if application was made before snapshot feature
        student = None
        if not p_snap:
            student = await db.db.students.find_one(id_query(app.get("student_id")))

        row = []
        for col in safe_columns:
            c_norm = col.lower().strip().replace(" ", "").replace("_", "").replace("-", "")
            
            if c_norm in ["studentid/rollnumber", "rollnumber", "rollno", "studentid", "registernumber"]:
                row.append(p_snap.get("roll_number") or (student.get("roll_number") if student else ""))
            elif c_norm in ["fullname", "name", "studentname", "candidatename"]:
                row.append(p_snap.get("name") or (student.get("name") if student else ""))
            elif c_norm in ["institutionalemail", "instituteemail", "domainemail", "collegeemail"]:
                row.append(p_snap.get("institutional_email") or (student.get("emails", {}).get("institute") if student else ""))
            elif c_norm in ["personalemail", "alternateemail", "email"]:
                row.append(p_snap.get("personal_email") or (student.get("emails", {}).get("personal") if student else ""))
            elif c_norm in ["mobilenumber", "mobile", "phone", "contactnumber"]:
                row.append(p_snap.get("mobile") or (student.get("mobile") if student else ""))
            elif c_norm in ["program", "degree", "course"]:
                row.append(p_snap.get("program") or (student.get("program_name") if student else ""))
            elif c_norm in ["department", "branch", "stream"]:
                row.append(p_snap.get("department") or (student.get("department_name") if student else ""))
            elif c_norm in ["class/section", "class", "section"]:
                row.append(p_snap.get("class_section") or (student.get("class_name") if student else ""))
            elif c_norm in ["yearofpassing", "passoutyear", "graduationyear"]:
                row.append(p_snap.get("passout_year") or (student.get("passout_year") if student else ""))
            elif c_norm in ["10thpercentage", "10thscore", "10th%"]:
                row.append(a_snap.get("tenth_percentage") or (student.get("education", {}).get("secondary", {}).get("normalized_percentage") if student else ""))
            elif c_norm in ["12th/diplomapercentage", "12thpercentage", "12thscore", "diplomapercentage"]:
                row.append(a_snap.get("twelfth_percentage") or (student.get("education", {}).get("higher_secondary_or_diploma", {}).get("normalized_percentage") if student else ""))
            elif c_norm in ["ugcgpa", "cgpa", "ugtotalcgpa"]:
                row.append(a_snap.get("ug_cgpa") or (student.get("cgpa") if student else 0.0))
            elif c_norm in ["ugpercentage", "ug%"]:
                row.append(a_snap.get("ug_percentage") or (student.get("education", {}).get("undergraduate", {}).get("normalized_percentage") if student else ""))
            elif c_norm in ["currentbacklogs", "activebacklogs", "backlogs"]:
                row.append(a_snap.get("current_backlogs") if a_snap.get("current_backlogs") is not None else (student.get("active_backlogs", 0) if student else 0))
            elif c_norm in ["skills", "technicalskills"]:
                s_list = p_snap.get("skills") or (student.get("skills") if student else [])
                row.append(", ".join(s_list) if isinstance(s_list, list) else str(s_list))
            elif c_norm in ["projects", "academicprojects"]:
                p_list = p_snap.get("projects") or (student.get("projects") if student else [])
                row.append(", ".join(p_list) if isinstance(p_list, list) else str(p_list))
            elif c_norm in ["internships", "workexperience"]:
                i_list = p_snap.get("internships") or (student.get("internships") if student else [])
                row.append(", ".join(i_list) if isinstance(i_list, list) else str(i_list))
            elif c_norm in ["certifications", "technicalcertifications"]:
                c_list = p_snap.get("certifications") or (student.get("certifications") if student else [])
                row.append(", ".join(c_list) if isinstance(c_list, list) else str(c_list))
            elif c_norm in ["linkedin", "linkedinurl"]:
                row.append(p_snap.get("links", {}).get("linkedin") or "")
            elif c_norm in ["github", "githuburl"]:
                row.append(p_snap.get("links", {}).get("github") or "")
            elif c_norm in ["portfolio", "portfoliourl"]:
                row.append(p_snap.get("links", {}).get("portfolio") or "")
            elif c_norm in ["resumeurl", "resumelink", "resume"]:
                row.append(app.get("resume_version") or p_snap.get("resume_url") or (student.get("resume_url") if student else ""))
            elif c_norm in ["applicationstatus", "status"]:
                row.append(app.get("status", "Applied"))
            elif c_norm in ["ctcoffered", "package"]:
                row.append(app.get("ctc_offered") or drive.get("package", ""))
            else:
                row.append("")
                
        ws.append(row)

    # Auto-fit column widths
    for column in ws.columns:
        max_len = max((len(str(cell.value or '')) for cell in column), default=10)
        col_letter = column[0].column_letter
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)

    await log_audit_event(
        action="hr_applications_exported",
        actor=current_user["email"],
        actor_role=current_user.get("role", "TPO"),
        entity_type="PlacementDrive",
        entity_id=drive_id,
        institution_id=institution_id,
        details={"record_count": len(apps), "columns": safe_columns}
    )

    company_slug = drive.get('company_name', 'Campus').replace(' ', '_')
    role_slug = drive.get('job_role', 'Hiring').replace(' ', '_')
    filename = f"{company_slug}_{role_slug}_HR_Shortlist.xlsx"

    return StreamingResponse(
        file_stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

# ── Seminar / Event Scheduling ──────────────────────────────────────────────
class EventCreatePayload(BaseModel):
    title: str
    speaker: str
    date_time: str
    venue: str
    description: Optional[str] = None
    target_type: str = "All"
    target_id: Optional[str] = None

@router.post("/events/create", status_code=status.HTTP_201_CREATED, dependencies=[Depends(RoleChecker(["TPO"]))])
async def create_event(
    payload: EventCreatePayload,
    current_user: dict = Depends(get_current_user),
    institution_id: str = Depends(get_tenant_id)
):
    event_doc = {
        "institution_id": institution_id,
        "title": payload.title.strip(),
        "speaker": payload.speaker.strip(),
        "date_time": payload.date_time.strip(),
        "venue": payload.venue.strip(),
        "description": payload.description.strip() if payload.description else "",
        "target_type": payload.target_type,
        "target_id": payload.target_id,
        "created_by": current_user["user_id"],
        "created_at": datetime.now(timezone.utc).isoformat()
    }

    res = await db.db.events.insert_one(event_doc)
    event_id = str(res.inserted_id)

    return {
        "message": "Seminar scheduled successfully!",
        "event_id": event_id
    }

@router.get("/events", dependencies=[Depends(RoleChecker(["TPO", "Dean"]))])
async def list_events(
    current_user: dict = Depends(get_current_user),
    institution_id: str = Depends(get_tenant_id)
):
    events_cursor = db.db.events.find({"institution_id": institution_id}).sort("created_at", -1)
    events = await events_cursor.to_list(length=100)
    result = []
    for e in events:
        result.append({
            "id": str(e["_id"]),
            "title": e.get("title"),
            "speaker": e.get("speaker"),
            "date_time": e.get("date_time"),
            "venue": e.get("venue"),
            "description": e.get("description"),
            "target_type": e.get("target_type", "All"),
            "target_id": e.get("target_id"),
            "created_at": e.get("created_at")
        })
    return result

# ── TPO Faculty Management ───────────────────────────────────────────────────
@router.get("/faculty", dependencies=[Depends(RoleChecker(["TPO", "Dean"]))])
async def list_tpo_faculty(
    current_user: dict = Depends(get_current_user),
    institution_id: str = Depends(get_tenant_id),
):
    faculty_cursor = db.db.users.find({
        "institution_id": institution_id,
        "role": "Faculty",
        "is_active": {"$ne": False}
    }, {"_id": 1, "name": 1, "email": 1, "role": 1, "phone": 1, "class_id": 1, "department_id": 1, "created_at": 1})

    faculty_list = await faculty_cursor.to_list(length=500)
    for f in faculty_list:
        f["id"] = str(f["_id"])
        del f["_id"]

    return faculty_list

@router.delete("/faculty/{user_id}", dependencies=[Depends(RoleChecker(["TPO", "Dean"]))])
async def remove_tpo_faculty(
    user_id: str,
    current_user: dict = Depends(get_current_user),
    institution_id: str = Depends(get_tenant_id),
):
    faculty = await db.db.users.find_one(id_query(user_id))
    if not faculty or faculty.get("role") != "Faculty":
        raise HTTPException(status_code=404, detail="Faculty coordinator not found.")

    # Remove faculty coordinator assignment from any mapped classes
    await db.db.classes.update_many(
        {
            "$or": [
                {"faculty_coordinator_id": user_id},
                {"faculty_coordinator_id": parse_id(user_id)},
                {"faculty_coordinator_id": str(faculty.get("_id"))}
            ]
        },
        {"$unset": {"faculty_coordinator_id": ""}}
    )

    # Remove user record completely so email can be reused and no stale records remain
    await db.db.users.delete_one(id_query(user_id))

    await log_audit_event(
        action="faculty_coordinator_removed",
        actor=current_user.get("email", ""),
        actor_role=current_user.get("role", "TPO"),
        entity_type="User",
        entity_id=user_id,
        institution_id=institution_id,
        details={"faculty_name": faculty.get("name"), "faculty_email": faculty.get("email")}
    )

    return {"message": f"Faculty Coordinator '{faculty.get('name')}' removed successfully."}

